from django.db import transaction
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib.auth import login as auth_login, authenticate, logout
from django.contrib.auth.forms import AuthenticationForm
from django.contrib.auth.models import User
from django.contrib import messages
from django.core.files.base import ContentFile
from django.core.mail import EmailMultiAlternatives
from django.utils.html import strip_tags
from django.http import JsonResponse, HttpResponse
from datetime import datetime
from docx import Document
import json
import re
import uuid
import threading
from .models import Evento, ImagenEvento
from .forms import EventoForm
import zipfile
import io
import random
from django.templatetags.static import static



def index(request):
    # 1. ¿El usuario intentó iniciar sesión desde el modal del inicio?
    if request.method == 'POST' and 'login_submit' in request.POST:
        form = AuthenticationForm(request, data=request.POST)
        if form.is_valid():
            usuario = form.get_user()
            auth_login(request, usuario)
            # ¡Ya entró! Recargamos la página de inicio para que aparezcan las cosas ocultas
            return redirect('inicio') 
        else:
            messages.error(request, 'Usuario o contraseña incorrectos.')
    
    # Si no es POST, preparamos un formulario de login vacío
    login_form = AuthenticationForm()

    # Buscamos los 3 eventos más recientes (ordenados por fecha descendente)
    ultimas_noticias = Evento.objects.all().order_by('-fecha', '-id')[:3]
    
    # --- GALERÍA INSTITUCIONAL DINÁMICA CON FOTOS DE BOLETINES ---
    pool_fotos = []
    vistos = set()

    # 1. Portadas de los boletines más recientes
    boletines_con_portada = Evento.objects.filter(
        imagen_portada__isnull=False
    ).exclude(imagen_portada='').order_by('-fecha', '-id')[:10]

    for ev in boletines_con_portada:
        if ev.imagen_portada and ev.imagen_portada.name:
            try:
                url = ev.imagen_portada.url
                if url and url not in vistos:
                    vistos.add(url)
                    pool_fotos.append({
                        'url': url,
                        'titulo': ev.titulo or 'Evento Institucional',
                        'fecha': ev.fecha.strftime('%d/%m/%Y') if ev.fecha else '',
                        'evento_id': ev.id
                    })
            except Exception:
                pass

    # 2. Fotos de galerías asociadas a boletines recientes
    fotos_galeria = ImagenEvento.objects.select_related('evento').exclude(
        imagen=''
    ).exclude(imagen__isnull=True).order_by('-evento__fecha', '-id')[:20]

    for img_obj in fotos_galeria:
        if img_obj.imagen and img_obj.imagen.name:
            try:
                url = img_obj.imagen.url
                if url and url not in vistos:
                    vistos.add(url)
                    pool_fotos.append({
                        'url': url,
                        'titulo': (img_obj.evento.titulo if img_obj.evento else None) or 'Galería Institucional',
                        'fecha': (img_obj.evento.fecha.strftime('%d/%m/%Y') if img_obj.evento and img_obj.evento.fecha else ''),
                        'evento_id': img_obj.evento.id if img_obj.evento else None
                    })
            except Exception:
                pass

    # 3. Tomar las 5 fotos más recientes de los boletines publicados
    galeria_fotos = pool_fotos[:5]

    # Únicamente si hay menos de 5 fotos registradas, rellenar con respaldos institucionales
    if len(galeria_fotos) < 5:
        fallbacks = [
            {'url': static('imagenes/galeria_foto_0.jpeg'), 'titulo': 'Evento Institucional'},
            {'url': static('imagenes/galeria_foto_2.jpeg'), 'titulo': 'Obra Pública Municipal'},
            {'url': static('imagenes/galeria_foto_3.jpeg'), 'titulo': 'Atención Ciudadana'},
            {'url': static('imagenes/galeria_foto_8.jpeg'), 'titulo': 'Panorámica de Atizapán'},
            {'url': static('imagenes/galeria_foto_5.jpeg'), 'titulo': 'Reunión Institucional'},
        ]
        urls_actuales = {f['url'] for f in galeria_fotos}
        for fb in fallbacks:
            if len(galeria_fotos) >= 5:
                break
            if fb['url'] not in urls_actuales:
                galeria_fotos.append(fb)
                urls_actuales.add(fb['url'])

    # Mandamos todo al template
    return render(request, 'index.html', {
        'ultimas_noticias': ultimas_noticias,
        'login_form': login_form,
        'galeria_fotos': galeria_fotos
    })

    
def calendario(request):
    eventos_db = Evento.objects.all()
    lista_eventos = []
    
    for ev in eventos_db:
        fotos_extra_urls = [img.imagen.url for img in ev.galeria.all()]
        
        lista_eventos.append({
            'id': ev.id, # <--- ¡AÑADE ESTA LÍNEA!
            'date': ev.fecha.strftime('%Y-%m-%d') if ev.fecha else '', 
            'time': ev.hora.strftime('%I:%M %p') if ev.hora else '',
            'title': ev.titulo if ev.titulo else 'Evento sin título',
            'tag': ev.ubicacion if ev.ubicacion else 'Ubicación pendiente',
            'desc': ev.descripcion if ev.descripcion else 'Sin descripción.',
            'cobertura': ev.cobertura if ev.cobertura else '',
            'imagen': ev.imagen_portada.url if ev.imagen_portada else None,
            'documento_url': ev.documento.url if ev.documento else None,
            'galeria': fotos_extra_urls,
            'contenido': ev.contenido if ev.contenido else ev.descripcion,
        })
    
    eventos_json = json.dumps(lista_eventos)
    return render(request, 'calendario.html', {'eventos_json': eventos_json})

@login_required(login_url='inicio') # Lo mandamos a inicio si no está logueado
def crear_evento(request):
    # --- NUEVA REGLA DE SEGURIDAD ---
    # Si el usuario NO es super admin, lo pateamos al inicio con un mensaje de error
    if not request.user.is_superuser:
        messages.error(request, 'Acceso denegado. Solo los administradores pueden subir noticias.')
        return redirect('inicio')
    # ---------------------------------

    if request.method == 'POST':
        form = EventoForm(request.POST, request.FILES)
        fotos_multiples = request.FILES.getlist('fotos_extra') 
        
        if form.is_valid():
            evento = form.save(commit=False)
            
            if 'documento' in request.FILES:
                try:
                    doc = Document(request.FILES['documento'])
                    parrafos = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                    if len(parrafos) >= 6:
                        if not evento.titulo: evento.titulo = parrafos[0] 
                        if not evento.descripcion: evento.descripcion = parrafos[5] 
                        if not evento.fecha: evento.fecha = datetime.today().date()
                    
                    imagenes_docx = []
                    for rel in doc.part.rels.values():
                        if "image" in rel.target_ref:
                            img_part = rel.target_part
                            img_bytes = img_part.blob
                            ext = img_part.content_type.split('/')[-1]
                            imagenes_docx.append({'bytes': img_bytes, 'ext': ext})
                    
                    if not evento.imagen_portada and imagenes_docx:
                        img_data = imagenes_docx.pop(0)
                        nombre_archivo = f"portada_auto_{datetime.now().timestamp()}.{img_data['ext']}"
                        evento.imagen_portada.save(nombre_archivo, ContentFile(img_data['bytes']), save=False)
                        
                except Exception as e:
                    print("Error procesando el Word:", e)
            
            evento.save()
            
            if 'documento' in request.FILES:
                for i, img_data in enumerate(imagenes_docx):
                    nueva_img = ImagenEvento(evento=evento)
                    nombre_archivo = f"galeria_auto_{evento.id}_{i}.{img_data['ext']}"
                    nueva_img.imagen.save(nombre_archivo, ContentFile(img_data['bytes']))

            for foto in fotos_multiples:
                ImagenEvento.objects.create(evento=evento, imagen=foto)
                
            return redirect('calendario')
    else:
        form = EventoForm()
    
    return render(request, 'crear_evento.html', {'form': form})

def login_view(request):
    if request.method == 'POST':
        form = AuthenticationForm(request, data=request.POST)
        if form.is_valid():
            usuario = form.get_user()
            auth_login(request, usuario)
            return redirect('crear_evento')
    else:
        form = AuthenticationForm()
    return render(request, 'login.html', {'form': form})

def logout_view(request):
    logout(request)
    return redirect('inicio')

def formatear_fecha_espanol(texto_fecha):
    meses = {'enero': '01', 'febrero': '02', 'marzo': '03', 'abril': '04', 'mayo': '05', 'junio': '06', 
             'julio': '07', 'agosto': '08', 'septiembre': '09', 'octubre': '10', 'noviembre': '11', 'diciembre': '12'}
    match = re.search(r'(\d{1,2})\s+de\s+([a-z]+)\s+de\s+(\d{4})', texto_fecha.lower())
    if match:
        dia = match.group(1).zfill(2)
        mes = meses.get(match.group(2), '01')
        anio = match.group(3)
        return f"{anio}-{mes}-{dia}"
    return ""

def registro_medio(request):
    if request.method == 'POST':
        nombre = request.POST.get('Nombre')
        medio = request.POST.get('Nombre_del_medio')
        tipo = request.POST.get('Tipo_de_medio')
        correo = request.POST.get('Correo')
        telefono = request.POST.get('Telefono')

        if User.objects.filter(username=correo).exists():
            messages.error(request, 'Este correo ya está registrado.')
            return redirect('inicio')

        password_temporal = str(uuid.uuid4())[:8]
        nuevo_usuario = User.objects.create_user(
            username=correo,
            email=correo,
            password=password_temporal,
            first_name=nombre,
            is_active=False
        )

        dominio = request.get_host()
        link_si = f"http://{dominio}/activar-cuenta/{nuevo_usuario.id}/"
        link_no = f"http://{dominio}/rechazar-cuenta/{nuevo_usuario.id}/"

        subject = f'Nueva Solicitud de Prensa: {medio}'
        html_content = f"""
        <div style="font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif; color: #333333; max-width: 600px; margin: 0 auto; background-color: #ffffff; border-radius: 8px; overflow: hidden; box-shadow: 0 4px 15px rgba(0,0,0,0.1); border: 1px solid #e1e1e1;">
            
            <div style="background: linear-gradient(135deg, #0f1f3d 0%, #1a365d 100%); padding: 30px 20px; text-align: center; border-bottom: 4px solid #cc9933;">
                <h1 style="color: #ffffff; margin: 0; font-size: 24px; font-weight: 600; letter-spacing: 1px;">Solicitud de Acceso</h1>
                <p style="color: #e2e8f0; margin: 10px 0 0 0; font-size: 15px;">Sala de Prensa Virtual</p>
            </div>
            
            <div style="padding: 35px 30px; background-color: #f8fafc;">
                <p style="font-size: 16px; line-height: 1.6; margin-top: 0; color: #475569;">Estimada licenciada Veronica,</p>
                <p style="font-size: 16px; line-height: 1.6; color: #475569;">Un nuevo representante de medios ha solicitado acceso a la plataforma. Por seguridad, su cuenta se encuentra <strong style="color: #dc2626;">inactiva</strong> hasta su revisión.</p>
                
                <div style="background-color: #ffffff; border-radius: 6px; padding: 25px; margin: 25px 0; border: 1px solid #e2e8f0; box-shadow: 0 2px 4px rgba(0,0,0,0.02);">
                    <h3 style="color: #0f1f3d; margin-top: 0; margin-bottom: 20px; font-size: 18px; border-bottom: 2px solid #f1f5f9; padding-bottom: 10px;">Datos del Solicitante</h3>
                    
                    <table style="width: 100%; border-collapse: collapse;">
                        <tr>
                            <td style="padding: 8px 0; color: #64748b; width: 35%;"><strong> Nombre:</strong></td>
                            <td style="padding: 8px 0; color: #1e293b; font-weight: 500;">{nombre}</td>
                        </tr>
                        <tr>
                            <td style="padding: 8px 0; color: #64748b;"><strong> Medio:</strong></td>
                            <td style="padding: 8px 0; color: #1e293b; font-weight: 500;">{medio}</td>
                        </tr>
                        <tr>
                            <td style="padding: 8px 0; color: #64748b;"><strong> Tipo:</strong></td>
                            <td style="padding: 8px 0; color: #1e293b; font-weight: 500;">{tipo}</td>
                        </tr>
                        <tr>
                            <td style="padding: 8px 0; color: #64748b;"><strong> Correo:</strong></td>
                            <td style="padding: 8px 0; color: #cc9933; font-weight: 500;">
                                <a href="mailto:{correo}" style="color: #cc9933; text-decoration: none;">{correo}</a>
                            </td>
                        </tr>
                        <tr>
                            <td style="padding: 8px 0; color: #64748b;"><strong> Teléfono:</strong></td>
                            <td style="padding: 8px 0; color: #1e293b; font-weight: 500;">{telefono}</td>
                        </tr>
                    </table>
                </div>
                
                <div style="text-align: center; margin-top: 35px; padding-top: 25px; border-top: 1px solid #e2e8f0;">
                    <p style="font-size: 16px; color: #334155; margin-bottom: 20px; font-weight: 600;">¿Deseas autorizar el acceso a este usuario?</p>
                    
                    <table width="100%" border="0" cellspacing="0" cellpadding="0">
                        <tr>
                            <td align="center">
                                <a href="{link_si}" style="background-color: #10b981; color: #ffffff; text-decoration: none; padding: 14px 28px; border-radius: 6px; font-weight: bold; font-size: 15px; display: inline-block; margin-right: 15px; border: 1px solid #059669; box-shadow: 0 2px 4px rgba(16, 185, 129, 0.2);">
                                     SÍ, ACTIVAR
                                </a>
                                <a href="{link_no}" style="background-color: #ef4444; color: #ffffff; text-decoration: none; padding: 14px 28px; border-radius: 6px; font-weight: bold; font-size: 15px; display: inline-block; border: 1px solid #b91c1c; box-shadow: 0 2px 4px rgba(239, 68, 68, 0.2);">
                                     NO AUTORIZAR
                                </a>
                            </td>
                        </tr>
                    </table>
                </div>
            </div>
            
            <div style="background-color: #f1f5f9; padding: 15px; text-align: center; font-size: 12px; color: #94a3b8; border-top: 1px solid #e2e8f0;">
                <p style="margin: 0;">Este es un correo automático del sistema de administración.</p>
                <p style="margin: 5px 0 0 0;">Gobierno de Atizapán de Zaragoza &copy; {datetime.now().year}</p>
            </div>
        </div>
        """
        
        text_content = strip_tags(html_content)
        from_email = 'sub.medios@atizapan.gob.mx'
        to = ['sub.medios@atizapan.gob.mx']

        try:
            threading.Thread(
                target=enviar_correo_background, 
                args=(subject, text_content, html_content, from_email, to[0])
            ).start()
        except Exception as e:
            print("Error al iniciar hilo de correo:", e)

        messages.success(request, '¡Tu solicitud fue enviada con éxito! Pronto revisaremos tus datos.')
        return redirect('inicio')
        
    return redirect('inicio')

def activar_cuenta(request, user_id):
    try:
        usuario = User.objects.get(id=user_id)
        
        if usuario.is_active:
            return HttpResponse("<h2 style='font-family: Arial; text-align: center; margin-top: 50px;'>Esta cuenta ya había sido activada previamente.</h2>")

        password_definitiva = str(uuid.uuid4())[:8]
        usuario.set_password(password_definitiva) 
        usuario.is_active = True
        usuario.is_staff = True
        usuario.save()

        subject = '¡Bienvenido a la Sala de Prensa de Atizapán de Zaragoza! '
        html_correo = f"""
        <div style="font-family: 'Montserrat', Arial, sans-serif; color: #333; max-width: 600px; margin: auto; border: 1px solid #ddd; border-radius: 10px; overflow: hidden;">
            <div style="background-color: #0f1f3d; padding: 30px; text-align: center; color: #fff;">
                <h2 style="margin: 0; font-family: 'Playfair Display', serif; font-size: 24px;">¡Bienvenido!</h2>
                <p style="margin-top: 10px; font-size: 16px;">Tu cuenta ha sido autorizada</p>
            </div>
            <div style="padding: 30px; background-color: #f9f9f9;">
                <p style="font-size: 16px;">Hola <strong>{usuario.first_name}</strong>,</p>
                <p style="font-size: 15px; line-height: 1.6;">El Gobierno de Atizapán de Zaragoza te da la más cordial bienvenida a nuestra <strong>Sala de Prensa Virtual</strong>.</p>
                <p style="font-size: 15px; line-height: 1.6;">A partir de este momento, puedes acceder a material exclusivo, galerías y recursos para medios.</p>
                
                <div style="background-color: #fff; padding: 20px; border-radius: 8px; border-left: 4px solid #cc9933; margin: 25px 0;">
                    <h3 style="margin-top: 0; color: #0f1f3d; font-family: 'Playfair Display', serif;">Tus Credenciales de Acceso:</h3>
                    <p style="margin: 8px 0; font-size: 16px;"><strong>Correo:</strong> {usuario.email}</p>
                    <p style="margin: 8px 0; font-size: 16px;"><strong>Contraseña:</strong> {password_definitiva}</p>
                </div>
                
                <div style="text-align: center; margin-top: 35px;">
                    <a href="http://{request.get_host()}/" style="background-color: #cc9933; color: white; padding: 14px 30px; text-decoration: none; border-radius: 5px; font-weight: bold; display: inline-block; text-transform: uppercase; letter-spacing: 1px;">Ingresar a la Sala</a>
                </div>
            </div>
        </div>
        """
        
        threading.Thread(
            target=enviar_correo_background, 
            args=(subject, strip_tags(html_correo), html_correo, 'sub.medios@atizapan.gob.mx', usuario.email)
        ).start()

        return render(request, 'cuenta_activada.html', {'email_usuario': usuario.email})

    except User.DoesNotExist:
        return HttpResponse("<h2 style='font-family: Arial; text-align: center; margin-top: 50px; color: red;'>Error: Ese usuario no existe.</h2>")

@login_required(login_url='login')
def extraer_datos_word(request):
    if request.method == 'POST' and request.FILES.get('documento'):
        try:
            doc = Document(request.FILES['documento'])
            parrafos = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            
            titulo = parrafos[0] if len(parrafos) > 0 else ""
            fecha_texto = parrafos[2] if len(parrafos) > 2 else ""
            fecha_lista = formatear_fecha_espanol(fecha_texto)
            descripcion = parrafos[5] if len(parrafos) > 5 else ""
            
            contenido_completo = "\n\n".join(parrafos[5:]) if len(parrafos) > 5 else ""

            return JsonResponse({
                'success': True,
                'titulo': titulo,
                'fecha': fecha_lista,
                'descripcion': descripcion,
                'contenido': contenido_completo
            })
        except Exception as e:
            return JsonResponse({'success': False, 'error': str(e)})
            
    return JsonResponse({'success': False, 'error': 'No se envió archivo'})

def enviar_correo_background(subject, text_content, html_content, from_email, to_email):
    try:
        msg = EmailMultiAlternatives(subject, text_content, from_email, [to_email])
        msg.attach_alternative(html_content, "text/html")
        msg.send()
        print(f"✅ ¡ÉXITO! Correo enviado correctamente a: {to_email}")
    except Exception as e:
        print(f"❌ ERROR AL ENVIAR CORREO a {to_email}. Detalle: {e}")


@login_required(login_url='inicio')
def descargar_pack_medios(request, evento_id):
    evento = get_object_or_404(Evento, id=evento_id)
    
    # Creamos un archivo ZIP en la memoria (sin guardarlo en el disco duro)
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
        
        # 1. Empaquetar el documento Word (si existe)
        if evento.documento and evento.documento.name:
            file_name = evento.documento.name.split('/')[-1]
            zip_file.writestr(f"Boletin_Oficial/{file_name}", evento.documento.read())
            
        # 2. Empaquetar la imagen de portada (si existe)
        if evento.imagen_portada and evento.imagen_portada.name:
            img_name = evento.imagen_portada.name.split('/')[-1]
            zip_file.writestr(f"Fotografias/portada_{img_name}", evento.imagen_portada.read())
            
        # 3. Empaquetar TODA la galería de fotos
        for i, img_obj in enumerate(evento.galeria.all()):
            if img_obj.imagen and img_obj.imagen.name:
                gal_name = img_obj.imagen.name.split('/')[-1]
                zip_file.writestr(f"Fotografias/galeria_{i}_{gal_name}", img_obj.imagen.read())
        
        # 4. EXTRA COOL: Creamos un archivo de texto con el boletín por si no tienen Word
        texto_nota = f"TÍTULO: {evento.titulo}\n\nFECHA: {evento.fecha}\n\nUBICACIÓN: {evento.ubicacion}\n\nDESCRIPCIÓN:\n{evento.descripcion}\n\nCONTENIDO COMPLETO:\n{evento.contenido}"
        zip_file.writestr("Texto_Noticia.txt", texto_nota.encode('utf-8'))

    # Preparamos el ZIP para enviarlo como descarga al navegador
    zip_buffer.seek(0)
    response = HttpResponse(zip_buffer, content_type='application/zip')
    # Le ponemos un nombre chulo al archivo ZIP
    nombre_zip = f'Pack_Prensa_Atizapan_Noticia_{evento.id}.zip'
    response['Content-Disposition'] = f'attachment; filename="{nombre_zip}"'
    
    return response

@login_required(login_url='inicio')
def galeria_recursos(request):
    # Traemos todos los eventos ordenados por fecha (los más nuevos primero)
    eventos_db = Evento.objects.all().order_by('-fecha')
    lista_eventos = []
    
    for ev in eventos_db:
        fotos_extra_urls = [img.imagen.url for img in ev.galeria.all()]
        
        lista_eventos.append({
            'id': ev.id,
            'date': ev.fecha.strftime('%Y-%m-%d') if ev.fecha else '',
            'title': ev.titulo if ev.titulo else 'Evento sin título',
            'tag': ev.ubicacion if ev.ubicacion else 'Atizapán',
            'desc': ev.descripcion if ev.descripcion else 'Sin descripción.',
            'cobertura': ev.cobertura if ev.cobertura else 'General',
            'imagen': ev.imagen_portada.url if ev.imagen_portada else None,
            'documento_url': ev.documento.url if ev.documento else None,
            'galeria': fotos_extra_urls,
            'contenido': ev.contenido if ev.contenido else ev.descripcion,
        })
    
    eventos_json = json.dumps(lista_eventos)
    # Mandamos la lista a nuestro nuevo archivo HTML
    return render(request, 'recursos.html', {'eventos_json': eventos_json})


def procesar_documento_boletin(file_obj, default_ubicacion='Atizapán de Zaragoza', default_cobertura='General'):
    """
    Extrae de forma robusta e inteligente los datos de un archivo Word (.docx):
    - Título (primer párrafo sustancial, limpiado y recortado para no exceder max_length=500)
    - Fecha (detectada por regex de fechas en español en los párrafos o fallback a hoy)
    - Descripción (párrafo introductorio de la nota)
    - Contenido (cuerpo completo de la nota estructurado)
    - Imágenes incrustadas (primera como portada, el resto como galería en ImagenEvento)
    - Guarda el archivo Word original adjunto al Evento
    """
    doc = Document(file_obj)
    
    # 1. Párrafos con texto no vacío
    parrafos = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    if not parrafos:
        raise ValueError("El documento Word no contiene texto o está vacío.")

    # 2. Extracción inteligente de Título
    titulo = parrafos[0]
    idx_inicio_cuerpo = 1
    
    # Comprobar si parrafos[0] es solo encabezado ("BOLETÍN 123", "COMUNICADO", etc.)
    if len(parrafos) > 1 and re.match(r'^(bolet[ií]n|comunicado|prensa)\b', parrafos[0].lower()):
        posible_titulo = f"{parrafos[0]} - {parrafos[1]}"
        titulo = posible_titulo if len(posible_titulo) <= 450 else parrafos[1]
        idx_inicio_cuerpo = 2
    
    # Limitar longitud para evitar error con max_length=500 de Django
    if len(titulo) > 480:
        titulo = titulo[:477] + "..."

    # 3. Extracción de Fecha mediante regex en español en los primeros párrafos
    fecha_final = datetime.today().date()
    meses = {
        'enero': 1, 'febrero': 2, 'marzo': 3, 'abril': 4, 'mayo': 5, 'junio': 6,
        'julio': 7, 'agosto': 8, 'septiembre': 9, 'octubre': 10, 'noviembre': 11, 'diciembre': 12
    }
    date_regex = re.compile(r'(\d{1,2})\s+de\s+([a-záéíóúñ]+)\s+de\s+(\d{4})', re.IGNORECASE)
    
    for p in parrafos[:10]:
        m = date_regex.search(p.lower())
        if m:
            try:
                dia = int(m.group(1))
                nom_mes = m.group(2).lower()
                anio = int(m.group(3))
                if nom_mes in meses and 1 <= dia <= 31 and 2000 <= anio <= 2100:
                    fecha_final = datetime(anio, meses[nom_mes], dia).date()
                    break
            except Exception:
                pass

    # 4. Descripción y Contenido Completo
    cuerpo_parrafos = parrafos[idx_inicio_cuerpo:]
    # Descartar párrafos cortos que solo contengan fecha o firma
    cuerpo_parrafos_filtrados = [
        p for p in cuerpo_parrafos 
        if not (len(p) < 80 and date_regex.search(p.lower()) and not p.endswith('.'))
    ]
    
    if cuerpo_parrafos_filtrados:
        descripcion = cuerpo_parrafos_filtrados[0]
        contenido = "\n\n".join(cuerpo_parrafos_filtrados)
    else:
        descripcion = titulo
        contenido = "\n\n".join(parrafos)

    # 5. Extracción de imágenes incrustadas en el docx
    imagenes_docx = []
    try:
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                img_part = rel.target_part
                img_bytes = img_part.blob
                ext = img_part.content_type.split('/')[-1]
                if ext.lower() == 'jpeg':
                    ext = 'jpg'
                imagenes_docx.append({'bytes': img_bytes, 'ext': ext})
    except Exception as img_err:
        print("Aviso al extraer imágenes del docx:", img_err)

    # 6. Creación del Evento
    evento = Evento(
        titulo=titulo,
        fecha=fecha_final,
        ubicacion=default_ubicacion or "Atizapán de Zaragoza",
        cobertura=default_cobertura or "General",
        descripcion=descripcion,
        contenido=contenido
    )
    
    # Asignar la primera foto como portada
    if imagenes_docx:
        portada_data = imagenes_docx.pop(0)
        nombre_portada = f"portada_auto_{uuid.uuid4().hex[:8]}.{portada_data['ext']}"
        evento.imagen_portada.save(nombre_portada, ContentFile(portada_data['bytes']), save=False)

    # Guardar el archivo Word original
    try:
        file_obj.seek(0)
    except Exception:
        pass
    nombre_doc = getattr(file_obj, 'name', f"boletin_{uuid.uuid4().hex[:8]}.docx")
    evento.documento.save(nombre_doc, file_obj, save=False)

    evento.save()

    # Guardar el resto de fotos en ImagenEvento (galería del boletín)
    for i, img_data in enumerate(imagenes_docx):
        nombre_galeria = f"galeria_auto_{evento.id}_{i}_{uuid.uuid4().hex[:6]}.{img_data['ext']}"
        nueva_img = ImagenEvento(evento=evento)
        nueva_img.imagen.save(nombre_galeria, ContentFile(img_data['bytes']))

    total_imagenes = (1 if evento.imagen_portada else 0) + len(imagenes_docx)
    return evento, total_imagenes


@login_required(login_url='inicio')
def carga_masiva_boletines(request):
    """
    Renderiza la vista de administración para subir y procesar lotes de boletines en Word.
    """
    if not request.user.is_superuser:
        messages.error(request, 'Acceso denegado. Solo los administradores pueden acceder a la carga masiva.')
        return redirect('inicio')
    return render(request, 'carga_masiva.html')


@login_required(login_url='inicio')
def procesar_boletin_individual(request):
    """
    Endpoint AJAX para procesar un único archivo Word a la vez dentro de la cola de carga masiva.
    Garantiza aislamiento, previene timeouts de servidor y transacciones atómicas seguras.
    """
    if not request.user.is_superuser:
        return JsonResponse({'success': False, 'error': 'Acceso no autorizado.'}, status=403)
        
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Método HTTP no permitido.'}, status=405)

    archivo = request.FILES.get('documento') or request.FILES.get('file')
    if not archivo:
        return JsonResponse({'success': False, 'error': 'No se recibió ningún archivo.'}, status=400)

    nombre_archivo = archivo.name
    if not nombre_archivo.lower().endswith('.docx'):
        return JsonResponse({
            'success': False, 
            'nombre_archivo': nombre_archivo,
            'error': f'Formato no compatible en "{nombre_archivo}". Sube archivos en formato Word (.docx).'
        }, status=400)

    ubicacion = request.POST.get('ubicacion', 'Atizapán de Zaragoza').strip() or 'Atizapán de Zaragoza'
    cobertura = request.POST.get('cobertura', 'General').strip() or 'General'

    try:
        with transaction.atomic():
            evento, total_fotos = procesar_documento_boletin(
                file_obj=archivo,
                default_ubicacion=ubicacion,
                default_cobertura=cobertura
            )

        return JsonResponse({
            'success': True,
            'id': evento.id,
            'titulo': evento.titulo,
            'fecha': evento.fecha.strftime('%d/%m/%Y') if evento.fecha else '',
            'portada_url': evento.imagen_portada.url if evento.imagen_portada else None,
            'total_fotos': total_fotos,
            'nombre_archivo': nombre_archivo,
            'message': 'Boletín publicado exitosamente en Sala de Prensa'
        })
    except Exception as e:
        return JsonResponse({
            'success': False,
            'nombre_archivo': nombre_archivo,
            'error': f'Error al procesar "{nombre_archivo}": {str(e)}'
        }, status=500)